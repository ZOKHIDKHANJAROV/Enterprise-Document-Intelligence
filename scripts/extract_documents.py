"""Extract enterprise documents into a normalized JSONL corpus.

Supported:
- PDF (native text, optional OCR for scanned pages)
- DOCX
- PPTX
- XLSX / XLSM
- XLS (via pandas + xlrd)
- TXT / MD

The script keeps source documents outside Git. Each output record contains a
stable SHA-256 document_id and source metadata.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from tqdm import tqdm

SUPPORTED = {
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".xlsm",
    ".xls",
    ".txt",
    ".md",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf_native(path: Path) -> tuple[str, int]:
    reader = PdfReader(str(path))
    pages: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"[PAGE {index}]\n{text}")

    return clean_text("\n\n".join(pages)), len(reader.pages)


def extract_pdf_ocr(path: Path, languages: str) -> tuple[str, int]:
    """OCR PDF pages using PyMuPDF + pytesseract.

    Optional dependencies:
        pip install pymupdf pytesseract
    Tesseract itself must also be installed on the operating system.
    """
    try:
        import fitz
        import pytesseract
    except ImportError as exc:
        raise RuntimeError(
            "OCR requires pymupdf and pytesseract. "
            "Install them with: pip install pymupdf pytesseract"
        ) from exc

    document = fitz.open(str(path))
    pages: list[str] = []

    try:
        for index, page in enumerate(document, start=1):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = pixmap.tobytes("png")
            from PIL import Image
            from io import BytesIO

            text = pytesseract.image_to_string(
                Image.open(BytesIO(image)),
                lang=languages,
            )
            pages.append(f"[PAGE {index}]\n{text}")
    finally:
        document.close()

    return clean_text("\n\n".join(pages)), len(pages)


def extract_pdf(path: Path, use_ocr: bool, ocr_languages: str) -> tuple[str, int, str]:
    native_text, pages = extract_pdf_native(path)

    # A scanned PDF often has almost no extractable text. In OCR mode, OCR the
    # complete document when native extraction is below a conservative threshold.
    if use_ocr and (len(native_text.strip()) < max(100, pages * 40)):
        ocr_text, ocr_pages = extract_pdf_ocr(path, ocr_languages)
        if len(ocr_text) > len(native_text):
            return ocr_text, ocr_pages, "ocr"

    return native_text, pages, "native"


def extract_docx(path: Path) -> tuple[str, int]:
    document = Document(str(path))
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            blocks.append(value)

    for table_index, table in enumerate(document.tables, start=1):
        blocks.append(f"[TABLE {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            blocks.append(" | ".join(cells))

    return clean_text("\n".join(blocks)), len(document.paragraphs)


def extract_pptx(path: Path) -> tuple[str, int]:
    presentation = Presentation(str(path))
    blocks: list[str] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        blocks.append(f"[SLIDE {slide_index}]")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                blocks.append(shape.text.strip())

    return clean_text("\n".join(blocks)), len(presentation.slides)


def extract_xlsx(path: Path) -> tuple[str, int]:
    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )
    blocks: list[str] = []
    rows = 0

    try:
        for sheet in workbook.worksheets:
            blocks.append(f"[SHEET: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    blocks.append(" | ".join(values))
                    rows += 1
    finally:
        workbook.close()

    return clean_text("\n".join(blocks)), rows


def extract_xls(path: Path) -> tuple[str, int]:
    blocks: list[str] = []
    rows = 0

    workbook = pd.ExcelFile(path, engine="xlrd")
    try:
        for sheet_name in workbook.sheet_names:
            frame = pd.read_excel(
                workbook,
                sheet_name=sheet_name,
                header=None,
            )
            blocks.append(f"[SHEET: {sheet_name}]")
            for row in frame.itertuples(index=False, name=None):
                values = ["" if pd.isna(value) else str(value) for value in row]
                if any(value.strip() for value in values):
                    blocks.append(" | ".join(values))
                    rows += 1
    finally:
        workbook.close()

    return clean_text("\n".join(blocks)), rows


def extract_text_file(path: Path) -> tuple[str, int]:
    return clean_text(
        path.read_text(encoding="utf-8", errors="ignore")
    ), 1


def extract(
    path: Path,
    use_ocr: bool,
    ocr_languages: str,
) -> tuple[str, int, str]:
    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(path, use_ocr, ocr_languages)

    if extension == ".docx":
        text, units = extract_docx(path)
        return text, units, "native"

    if extension == ".pptx":
        text, units = extract_pptx(path)
        return text, units, "native"

    if extension in {".xlsx", ".xlsm"}:
        text, units = extract_xlsx(path)
        return text, units, "native"

    if extension == ".xls":
        text, units = extract_xls(path)
        return text, units, "native"

    if extension in {".txt", ".md"}:
        text, units = extract_text_file(path)
        return text, units, "native"

    raise ValueError(f"Unsupported extension: {extension}")


def iter_documents(root: Path) -> Iterable[Path]:
    for path in sorted(root.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED:
            yield path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/extracted/corpus.jsonl"),
    )
    parser.add_argument("--max-documents", type=int, default=None)
    parser.add_argument(
        "--ocr",
        action="store_true",
        help="OCR scanned PDFs when native text extraction is insufficient.",
    )
    parser.add_argument(
        "--ocr-languages",
        default="rus+eng",
        help="Tesseract language code, e.g. rus+eng or eng.",
    )
    args = parser.parse_args()

    documents = list(iter_documents(args.input_dir))
    if args.max_documents is not None:
        documents = documents[: args.max_documents]

    args.output.parent.mkdir(parents=True, exist_ok=True)

    stats = {
        "discovered": len(documents),
        "processed": 0,
        "empty": 0,
        "failed": 0,
        "ocr_documents": 0,
        "errors": [],
        "extensions": {},
    }

    with args.output.open("w", encoding="utf-8") as output:
        for path in tqdm(documents, desc="Extracting"):
            extension = path.suffix.lower()
            stats["extensions"][extension] = (
                stats["extensions"].get(extension, 0) + 1
            )

            try:
                text, units, extraction_method = extract(
                    path,
                    use_ocr=args.ocr,
                    ocr_languages=args.ocr_languages,
                )
                digest = sha256_file(path)

                record = {
                    "document_id": digest,
                    "file_name": path.name,
                    "relative_path": str(path.relative_to(args.input_dir)),
                    "extension": extension,
                    "extraction_method": extraction_method,
                    "text": text,
                    "text_length": len(text),
                    "units": units,
                    "sha256": digest,
                    "extracted_at": datetime.now(timezone.utc).isoformat(),
                }

                output.write(
                    json.dumps(record, ensure_ascii=False) + "\n"
                )
                stats["processed"] += 1

                if extraction_method == "ocr":
                    stats["ocr_documents"] += 1

                if not text:
                    stats["empty"] += 1

            except Exception as exc:
                stats["failed"] += 1
                stats["errors"].append(
                    {"file": str(path), "error": repr(exc)}
                )

    stats_path = args.output.with_suffix(".stats.json")
    stats_path.write_text(
        json.dumps(stats, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
